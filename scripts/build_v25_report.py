"""Author the branded technical report from the checked-in audit evidence.

Documentation-only dependencies: python-docx and matplotlib. Render and inspect
the resulting DOCX separately before release; authoring is not a layout check.
"""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT/'Documents'
ASSETS = OUT/'assets'
DATA = json.loads((OUT/'v25_synthetic_benchmark.json').read_text())
NAVY = '132D46'


def chart():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import numpy as np
    names = ['static_albedo','static_albedo_jitter','subpixel_caster',
             'subpixel_caster_jitter','subpixel_caster_high_noise','metre_caster']
    labels = ['Static\nalbedo','Static albedo\n0.25 px jitter','0.3 m caster',
              '0.3 m caster\n0.25 px jitter','0.3 m caster\nhigher noise','1 m caster']
    fig, ax = plt.subplots(figsize=(9,3.7),layout='constrained')
    for i,key in enumerate(names):
        values = [r['top_score'] for r in DATA['scenarios'][key]['trials']]
        color = '#B24643' if i < 2 else '#173C5B'
        ax.scatter(i+np.linspace(-.13,.13,len(values)),values,s=25,color=color,alpha=.7)
        ax.plot([i-.25,i+.25],[np.median(values)]*2,color=color,lw=2.4)
    ax.set_yscale('log')
    ax.set_ylabel('Highest ranked score  (log scale)',fontsize=11)
    ax.set_xticks(range(6),labels,fontsize=10)
    ax.spines[['top','right']].set_visible(False)
    ax.grid(axis='y',alpha=.18)
    ax.set_axisbelow(True)
    fig.savefig(ASSETS/'v25_synthetic_scores.png',dpi=220)
    plt.close(fig)


doc = Document()
section = doc.sections[0]
section.page_width,section.page_height = Inches(8.5),Inches(11)
section.top_margin,section.bottom_margin = Inches(.7),Inches(.65)
section.left_margin,section.right_margin = Inches(.8),Inches(.8)
section.header_distance,section.footer_distance = Inches(.25),Inches(.25)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','Header','Footer']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_after = Pt(7)
    for border in list(style._element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(31)
doc.styles['Title'].font.bold = True
doc.styles['Subtitle'].font.size = Pt(16)
doc.styles['Heading 1'].font.size = Pt(21)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Caption'].font.size = Pt(9)
for name in ['Header','Footer']:
    doc.styles[name].font.size = Pt(9)
    doc.styles[name].paragraph_format.space_after = Pt(0)
header = section.header.paragraphs[0]
header.text = 'HATI     HAZARD ASSESSMENT AND TERRAIN INTELLIGENCE'
footer = section.footer.paragraphs[0]
footer.text = 'Scientific audit and release     •     10 September 2026'
for tabs in list(doc.styles['Footer']._element.iter(qn('w:tabs'))):
    tabs.getparent().remove(tabs)
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.85),WD_ALIGN_PARAGRAPH.RIGHT)
footer.add_run('\t')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'),'PAGE')
footer._p.append(field)
doc.core_properties.title = 'HATI v2 5 Scientific audit and release'
doc.core_properties.subject = 'Experimental lunar polar subpixel relief sensing and adversarial software audit'
doc.core_properties.author = 'HATI project'
doc.core_properties.keywords = 'HATI, lunar polar, shadow roots, scientific audit, experimental'


def para(text,boldlead=None):
    p = doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold=True
        p.add_run(text[len(boldlead):])
    else:
        p.add_run(text)
    return p


def heading(text,level=2):
    return doc.add_heading(text,level)


def page(title):
    doc.add_page_break()
    heading(title,1)


def table(headers,rows,widths):
    t = doc.add_table(rows=1,cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col,width in zip(t.columns,widths):
        col.width = Inches(width)
    for cell,text in zip(t.rows[0].cells,headers):
        cell.text = text
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for cell,text in zip(cells,row):
            cell.text=str(text)
    for i,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd')
            shade.set(qn('w:fill'),NAVY if i==0 else ('F1F5F8' if i%2==0 else 'FFFFFF'))
            props.append(shade)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                edge=OxmlElement('w:'+side)
                for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]:
                    edge.set(qn('w:'+key),val)
                borders.append(edge)
            props.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                margin=OxmlElement('w:'+side)
                margin.set(qn('w:w'),'95')
                margin.set(qn('w:type'),'dxa')
                margins.append(margin)
            props.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1)
                p.paragraph_format.line_spacing=1.03
                for run in p.runs:
                    run.font.size=Pt(10)
                    if i==0:
                        run.bold=True
                        run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t


def link(label,url):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(5)
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'),rel)
    r=OxmlElement('w:r')
    prop=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),NAVY); prop.append(color)
    r.append(prop)
    text=OxmlElement('w:t'); text.text=label; r.append(text)
    h.append(r); p._p.append(h)


def equation(text):
    # Simple expressions authored as native Office Math, not raw LaTeX.
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    math=OxmlElement('m:oMath')
    run=OxmlElement('m:r')
    node=OxmlElement('m:t'); node.text=text; run.append(node)
    math.append(run); p._p.append(math)


chart()
p=doc.add_paragraph()
p.add_run().add_picture(str(ASSETS/'hati-logo.png'),width=Inches(1.8))
doc.add_paragraph('HATI v2 5\nScientific audit and release',style='Title')
doc.add_paragraph('Lunar polar subpixel relief sensing',style='Subtitle')
para('10 September 2026  •  Core version 2.5.2  •  Scientific development release')
para('HATI now has an experimental shadow-root detector that fits multiple illuminations together and explicitly removes stationary albedo. The accompanying audit repairs concrete errors in acquisition selection, calibration, registration, statistics and georeferencing. The physical approach remains promising; measured lunar subpixel sensitivity is still to be established.')
table(['Evidence level','Release position'],[
    ('Implemented','Audited ingestion and experimental joint-image likelihood search'),
    ('Verified offline','Six test suites and 72 independent-renderer synthetic trials'),
    ('Still to measure','Real sweep recovery, local registration accuracy and false discoveries'),
    ('Not established','Object dimensions, hazard probabilities or landing clearance')],[1.5,5.35])
para('For the HATI project team. This report explains what changed, which conclusions survive an adversarial review, and what must be measured on the GPU host before making a performance claim. No ISIS application or ingestion was run on the Windows audit machine.')

page('1 Findings that change the conclusion')
para('The principal defect was inferential: spatial recurrence was treated as evidence that shadows followed the Sun. Persistent dark markings also recur. A better segmentation threshold or shadow-root proposal cannot repair that conclusion on its own.')
heading('Stationary albedo passes the old test')
para('Eight identical 512 by 512 images contained a nine-pixel dark stripe and no moving feature. With azimuths spanning 88 degrees, the legacy voter generated 29 pixels with at least four votes. Its original spatial-shift control produced zero exceedances in 200 trials. Four unfavorable frames abstained; the favorable subset repeatedly selected the same stripe endpoint. This is an executable counterexample to albedo rejection.')
heading('Selection could include postlanding imagery')
para('The existing selection plan contains March 2025 products acquired after the March 6 event. Replaying the old selector chose M1496097780LE from March 8 as a primary frame. The actual eight-frame Linux manifest was unavailable, so contamination of that particular run is not established. Acquisition timestamps now pass a strict exclusive UTC cutoff before primary or sibling selection.')
heading('Registration could leave its search bound')
para('The coarse correlation correction had the wrong sign. Fine registration could then compensate without respecting the declared bound. A planted 60-pixel displacement was returned despite a 20-pixel maximum. The correction sign and both refinement and total-shift bounds are repaired. Pairwise closure remains a consistency check: a coherently moving illumination pattern can close while aligning the wrong thing.')
heading('Synthetic recovery overstated small object sensitivity')
para('The previous width renderer quantized a range of nominal fractional widths to the same integer brush. Corrected pixel integration exposed failures in old half-metre smoke tests. Those machinery checks now explicitly use one-metre casters; their success cannot preserve the old half-metre sensitivity claim. The new benchmark uses a separate tapered-caster renderer.')
heading('Secondary defects also changed results')
para('Arc fitting summed row and column errors before squaring, allowing cancellation: a reproduced near-zero residual should have been 5.66 pixels. Positive recovery silently used a narrower width limit than the real detector. Unknown DEM pixels could pass a slope gate. Output origins mixed pixel centres and corners. Each has a code correction and offline regression evidence.')

page('2 Physics and coordinate geometry')
heading('Projected bearings')
para('Adding site longitude to ground solar azimuth is correct for the existing south-polar stereographic map with zero central meridian. It is not a universal polar formula. The new helper projects a one-metre geodesic displacement and applies the inverse raster affine transform. This supports north and south polar maps and rotated grids. Tests include both poles; the real-data adapter still targets Athena.')
heading('Shadow length is conditional on the receiving surface')
equation('h = L (tan e + tan β)')
para('Here h is caster height, L is horizontal shadow length, e is solar elevation and β is receiving-plane slope in the down-Sun direction. A rising surface shortens the shadow. At four degrees of elevation, an unmodelled one-degree slope changes the height-length factor by about 25 percent. If the denominator is nonpositive, a finite intersection does not exist in this plane model; the implementation rejects it.')
para('The root detector accepts row and column receiving-plane gradients. The default flat plane is an explicit sensitivity assumption, not a terrain correction. Censored endpoints support a root ranking but cannot measure height. Width is varied independently from height; the old two-to-one aspect ratio is not treated as a physical law.')
heading('Subpixel sampling is not resolved morphology')
para('An unresolved caster can alter several shadow pixels, allowing detection through integrated evidence. That does not resolve its shape or identify height independently of width, darkness, slope and blur. The 0.9-metre output posting does not establish the native optical resolution. Finite-Sun penumbra, detector sampling, map resampling and uncertainty in alignment all affect the footprint of the signal.')
para('The implemented forward model integrates fractional pixels, approximates a finite solar disc with five strips and applies a Gaussian point-spread function. An additional Gaussian width represents an assumed registration blur. It does not remove shifted-albedo residuals or replace local alignment measurements. A constant plane and rectangular shadow remain approximations; resolved-terrain ray tracing is not delivered.')
heading('ISIS shape and emission geometry')
para('A spherical map projection does not prove that camera rays intersected a sphere. ISIS spiceinit can select a system shape surface. The audited chain saves complete cube labels so Kernels.ShapeModel can be inspected on Linux. An emission cap of 40 degrees is a selection heuristic; it does not certify negligible relief parallax. Local independent checkpoints and the actual shape surface are required before subpixel localization claims.')

page('3 Statistical replacement')
heading('A null model that retains the nuisance')
equation('Y = A + P − c T + ε')
para('The image stack Y contains a shared static image A, a separate brightness plane P in every frame, a moving shadow template T, a shared nonnegative contrast c and residual noise. Under the null, contrast is zero. Both hypotheses use the same nuisance projection and common observed pixels. The implementation removes the static image and brightness planes with an exact weighted projection.')
para('This is a specific additive normalized-radiance model. It rejects exactly stationary albedo under its assumptions, including frame brightness planes. Nonlinear photometry, topographic shading, cast-shadow occlusion, albedo-dependent shadow depth and registration errors can escape that model. The new statistic improves the question being asked; it does not make the physical nuisance model complete.')
heading('All frames contribute to one fit')
para('Directional darkening in the nuisance residual proposes locations. Each proposed root is then evaluated across every illumination using a bank of height, width and subpixel-position templates. Contrast is shared and bounded. Contradictory observations reduce the joint fit improvement; there is no per-candidate permission to discard those frames. Missing pixels are excluded from all frames, and insufficient common support is reported.')
para('The score is the square root of the constrained reduction in weighted residual sum of squares. It is a ranking statistic. The output also records per-frame improvements, template identifiability after nuisance removal, common support and whether the candidate budget was exhausted. It does not convert the score into a hazard probability.')
heading('Selection belongs inside calibration')
para('The optional Gaussian model check simulates the complete no-shadow stack and reruns proposals, location search and all templates on every draw. The observed maximum is ranked against the simulated maxima using the plus-one rule: add one to both the exceedance count and trial count. This includes the search selection under a fixed independent-Gaussian model with known noise and support.')
para('Real NAC data violate that simple model through correlation, median normalization, uncertain noise and registration. Consequently the resulting p-value is a model diagnostic, not a measured lunar false-positive rate. The small smoke check uses 19 trials, for a minimum attainable p-value of 0.05; its observed rank is 0.30. It is a reproducibility check, not validation of calibration.')
heading('Legacy controls retain a limited role')
para('Spatial shifts of completed vote maps now use periodic translations that preserve vote mass. Finite-trial and threshold-selection corrections improve reporting. The control still tests spatial correspondence rather than albedo rejection. A small spatial-shift tail probability is not evidence of boulders, and no-vote pixels are not safe terrain.')

page('4 Implemented root search')
table(['Stage','Implemented behavior'],[
    ('Input','Audited, gate-passing frames; measured geometry; positive median-normalized radiance'),
    ('Proposal','Directional steps in residuals after removing static albedo and brightness planes'),
    ('Local fit','Circular support of radius 6 pixels; render radius 12 pixels plus blur padding'),
    ('Templates','Independent height and width; fractional coverage; half-pixel root offsets'),
    ('Evidence','Shared bounded contrast; all-frame likelihood improvement; identifiability check'),
    ('Output','Ranked roots, configuration and provenance, censored endpoints and common support')],[1.05,5.8])
heading('The root support limits later shadow mergers')
para('The final fit uses only a circular region around the candidate origin. A shadow that merges with terrain outside this support does not alter the local residual fit. Nearby merged shadows, occluded roots and illumination-changing terrain can still interfere. The proposal stage has a finite budget, so a scene with many stronger structures may hide weaker roots. That budget and truncation status are part of the recorded result.')
heading('Unknown is a product state')
para('The pilot writes a candidate table, a complete run JSON and a common-support raster. Support indicates whether every frame observes a pixel; it is not a clearance mask. Template dimensions are named as template parameters. Endpoints outside fitting support are marked censored. No dense probability map is manufactured from sparse candidates.')
heading('Terrain descriptors and object evidence remain separate')
para('DEM slope and roughness describe terrain at their supported baselines. Shadow likelihood measures image evidence for relief below or near native sampling. A boulder and a crater can both threaten landing, so the decision product should combine footprint-scale constraints with relief evidence and uncertainty; it should not require an unreliable morphology label before reporting potential relief.')
para('This release adds an optional TRI descriptor sampled at a fixed physical radius and corrects tied AUC ranks. Switching TRI definitions requires new reference normalization. Historical massif-versus-mare AUC values measure site separation, not safe-versus-unsafe object classification. They cannot establish universal channel weights or validate a landing decision.')

page('5 Offline validation and stress results')
para('All six offline suites pass. They cover channels, footprint geometry, registration, legacy kinematics, 14 audit regressions and the new likelihood tests. No external imagery acquisition or ISIS executable was used. The synthetic stress benchmark adds 72 trials: 12 seeds in each of six scenarios.')
doc.add_picture(str(ASSETS/'v25_synthetic_scores.png'),width=Inches(6.85))
doc.add_paragraph('Each point is one synthetic scene. Bars are medians. Scores are model rankings, not field significances.',style='Caption')
names=[('Static albedo','static_albedo'),('Static albedo with jitter','static_albedo_jitter'),
       ('0.3 m caster','subpixel_caster'),('0.3 m caster with jitter','subpixel_caster_jitter'),
       ('0.3 m caster with higher noise','subpixel_caster_high_noise'),('1 m caster','metre_caster')]
rows=[]
for label,key in names:
    s=DATA['scenarios'][key]
    rows.append((label,f"{s['median_top_score']:.2f}",
                 'Not applicable' if s['localized_count'] is None else f"{s['localized_count']} of 12"))
table(['Synthetic scenario','Median top score','Root localized'],rows,[3.4,1.45,2.0])
para('The 0.3-metre caster is 0.6 metres wide on a 0.9-metre grid. Localization means the highest-ranked root lies within two pixels of the planted position; no score cutoff is used. These favorable scenes use seven Sun directions spanning 310 degrees, an independent tapered-caster renderer and an anisotropic PSF. They do not represent the available Athena sweep.')
para('Registration jitter of 0.25 pixels raises the median selected score for stationary stains from 4.33 to 19.88. This is a demonstrated false-candidate mechanism. Even 12 successes in 12 synthetic trials do not establish near-perfect field sensitivity: the one-sided exact 95-percent lower bound for that same synthetic experiment is only about 78 percent.')

page('6 Ingestion and reproducibility repairs')
heading('Acquisition and cache contracts')
para('Selection requires known, pre-cutoff acquisition times before building alternate-frame lists. Ambiguous catalog selection fails rather than guessing. Downloads are staged and checked against complete PDS record counts before atomic replacement. Projected caches require a matching processing contract, geometry coverage and retained labels; failed rebuilds cannot leave a stale success contract.')
heading('A complete recorded calibration chain')
para('The chain runs lronac2isis, spiceinit, lronaccal, lronacecho and cam2map, with site coverage and measured incidence/emission checks. The added NAC echo correction follows the LROC processing guidance. Unknown campt coverage no longer counts as success. Cube labels, site geometry, map hash, source URL, acquisition time and processing version support review of the Linux run.')
heading('Alignment gates fail closed')
para('The coarse shift sign, refinement bound and final displacement bound are corrected. A manifest must have enough surviving frames, finite closure and passing gates; a failed ingest exits nonzero. Pairwise measurements are retained. Small residuals against the same reference are insufficient to prove correctness, and closure is not an independent local uncertainty bound. Those limitations remain explicit in the new runner.')
heading('Numerical and output corrections')
table(['Correction','Why it matters'],[
    ('Arc RMS','Squares vector components before summation; prevents error cancellation'),
    ('Injection width','Supersampled coverage replaces an integer brush; old recovery claims change'),
    ('Recovery configuration','Positive controls use the same configured width cap as real detection'),
    ('Connectivity','Diagonal components use the same eight-neighbour rule throughout'),
    ('Unknown DEM pixels','Masked support stays excluded from terrain clearance gates'),
    ('GeoTIFF transform','Reference corner transform and pixel-centre coordinates are kept distinct'),
    ('AUC ties','Equal scores receive half credit and no longer depend on input ordering')],[1.6,5.25])
para('The reference image and DEM labels were available locally and inspected. The actual eight-frame Linux cubes, geometry sidecars and closure records were unavailable. Therefore this report verifies code and local metadata, while the processing chain, frame yield and real scientific outputs still require the GPU-host run.')

page('7 Validation needed for a lunar result')
heading('First establish local registration and observability')
para('Measure held-out, terrain-fixed checkpoints near candidate roots, separated from the image textures used to estimate alignment. Quantify errors as a function of illumination, terrain and location. Verify the ISIS shape model and build resolved-terrain shadow/parallax predictions. Keep permanently shadowed, clipped, occluded and poorly registered areas explicitly unknown.')
heading('Separate calibration from independent truth')
para('Use multiple sites and solar sweeps, with calibration sites separated from test sites before selecting templates or thresholds. Build independently annotated object lists with adjudication and documented completeness. Human-reviewed Apollo 17 boulder catalogs, high-resolution local maps and published manual boulder annotations are starting points, not automatic subpixel ground truth for polar NAC imagery.')
para('Candidate resources include the Apollo 17 map archive, the Powell and colleagues manual boulder study and the BOULDERING annotation release. Verify their annotation protocol, effective image resolution, size completeness, coordinates and site coverage before using any as truth. Learned global boulder maps can suggest candidates but do not independently validate a detector under a no-trained-decision requirement.')
heading('A keystone measurement')
para('Report object-level precision, recall and false discoveries per surveyed area on held-out sites, stratified by height or width class, native resolution, incidence, emission, illumination diversity and terrain. Use one-to-one associations and spatial blocks for uncertainty intervals. Report non-observable area, missed objects and abstentions. Synthetic injection belongs before registration and native resampling to measure the full pipeline, alongside independent real annotations.')
para('A size-frequency distribution becomes defensible only after size-dependent completeness, false discoveries, association error and dimension uncertainty are estimated. Landing clearance additionally needs lander footprint limits, navigation error, reachable terrain and conservative treatment of missing observations. Neither a high scalar fusion score nor a low vote count supplies those requirements.')
heading('What can be published now')
para('The adversarial counterexample, repaired reproducible processing chain, explicit image-domain hypothesis test and independent-renderer stress characterization form a useful methods result. A claim of demonstrated subpixel lunar hazard sensing above 70 degrees requires independent polar data and external validation. The release provides the machinery to investigate that claim; it does not substitute favorable simulations for the missing measurement.')

page('8 Linux execution and evidence package')
heading('Run on the GPU host')
para('Use the existing Linux checkout and activate its ISIS environment. Pull the feat/v2.0-heatmap branch, install requirements-science.txt and run bash scripts/run_v25_wsl.sh. The checked-in WSL runbook contains the copyable terminal sequence. No GPU acceleration is implied: ISIS and this NumPy/SciPy search principally use CPU resources.')
para('The shell runner checks dependencies and all six test suites, queries pre-event imagery, ingests up to 24 selection bins and evaluates a central 512-pixel pilot. It stops on failures, including through logged pipes. Legacy projections rebuild under the audited processing contract while complete EDR downloads remain reusable.')
para('Three pilot runs assume registration blur of 0.25, 0.5 and 1 pixel. Noise sigma 0.03, optical PSF sigma 0.6 pixel and a flat receiving plane are also assumptions. Comparing their candidate stability is a sensitivity study, not a calibration. Replace these values only with documented local measurements or clearly identified additional scenarios.')
heading('Reviewable files')
table(['File or output','Purpose'],[
    ('AUDIT_2026-09-09.md','Detailed baseline findings with immutable code references and derivations'),
    ('V25_RELEASE.md','Implementation status, scientific scope and remaining research work'),
    ('WSL_AUDIT_RUN.md','Copyable Linux commands and interpretation of outputs'),
    ('v25_synthetic_benchmark.json','All 72 synthetic trials and the Gaussian full-search smoke check'),
    ('run.json and candidates.csv','Geometry, assumptions, configuration hash, frame evidence and ranked roots'),
    ('common_support.tif','Observed common area; neither observed nor unobserved means safe')],[2.35,4.5])
heading('Release acceptance and practical limits')
para('Software acceptance is the passing offline evidence and a reviewable, versioned implementation. Linux acceptance requires successful calibration, archive completeness, retained shape-model provenance and a passing ingest gate. Scientific acceptance requires independent local alignment and held-out truth. A gate pass alone does not establish the third level.')
para('The audit starts at commit 0fac80d2561915beffb9d8d04052bd035ca21486. This report accompanies the v2.5 release commit; the runner records the exact checked-out revision with every execution. The repository retains historical model-based material, clearly separated in the README from the deterministic v2.5 scientific path.')

page('9 Sources and evidence provenance')
para('The primary technical sources below support the processing and validation discussion. The numerical defect reproductions, implementation and synthetic results are original repository evidence, with executable tests and the full benchmark JSON. External resources are candidates for validation; their mere availability does not establish ground truth for this site.')
heading('ISIS and LROC processing')
link('LROC NAC Processing Guide — calibration chain and echo correction',
     'https://lroc.im-ldi.com/data/support/downloads/LROC_NAC_Processing_Guide.pdf')
link('USGS ISIS spiceinit — shape surface selection and kernels',
     'https://isis.astrogeology.usgs.gov/9.0.0/Application/presentation/Tabbed/spiceinit/spiceinit.html')
link('USGS ISIS campt — ground geometry and coordinate conventions',
     'https://isis.astrogeology.usgs.gov/9.0.0/Application/presentation/Tabbed/campt/campt.html')
link('USGS ISIS lronacecho — NAC echo correction',
     'https://isis.astrogeology.usgs.gov/9.0.0/Application/presentation/Tabbed/lronacecho/lronacecho.html')
link('USGS ISIS catlab — preserving cube labels',
     'https://isis.astrogeology.usgs.gov/9.0.0/Application/presentation/Tabbed/catlab/catlab.html')
heading('Statistics and independent validation resources')
link('Phipson and Smyth — finite Monte Carlo p values',
     'https://gksmyth.github.io/pubs/PermPValuesPreprint.pdf')
link('Apollo 17 Analyst Notebook — maps and contextual data',
     'https://an.rsl.wustl.edu/apollo/map/apMap.aspx?m=17')
link('Powell and colleagues — Apollo 17 boulder study',
     'https://agupubs.onlinelibrary.wiley.com/doi/10.1029/2022JE007532')
link('BOULDERING — published boulder annotation data',
     'https://zenodo.org/records/14250874')
heading('Repository evidence')
link('HATI repository — scientific branch and release history',
     'https://github.com/Cyrex567/HATI/tree/feat/v2.0-heatmap')
para('Audit baseline: 0fac80d2561915beffb9d8d04052bd035ca21486. The detailed Markdown audit links individual findings to this immutable revision. Release tests, raw synthetic trials, the branded logo and this report builder are retained in the repository so the document can be regenerated and inspected.')
para('Brand asset: existing HATI website logo, reused unchanged. Report preparation date: 10 September 2026. No real-data ingestion or GPU-host result is represented as completed here.')

path=OUT/'HATI_v25_Scientific_Audit_and_Release.docx'
doc.save(path)
print(path)
